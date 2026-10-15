"""Render DEPLOYMENT-GUIDE.md as a Word document.

Written for this one file, not as a general Markdown converter: it handles the
constructs the guide actually uses (ATX headings, pipe tables, fenced code,
bullets, bold/inline-code spans, blockquote-free prose) and would need work for
anything else.
"""

import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SRC, OUT = sys.argv[1], sys.argv[2]

MONO = "Consolas"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
CODE_BG = "F2F2F2"
CODE_FG = RGBColor(0x1A, 0x1A, 0x1A)

doc = Document()

# US Letter, 1" margins.
for s in doc.sections:
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(s, m, Inches(1))

base = doc.styles["Normal"]
base.font.name = "Calibri"
base.font.size = Pt(10.5)
base.paragraph_format.space_after = Pt(6)


def shade(el, fill):
    """Cell/paragraph background. python-docx has no API for either."""
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), fill)
    el.append(sh)


def inline(par, text):
    """Bold (**x**) and inline code (`x`); everything else literal."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = MONO
            r.font.size = Pt(9)
            r.font.color.rgb = ACCENT
        else:
            par.add_run(part)


def code_block(lines):
    """One shaded, monospaced paragraph per line, kept together."""
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(6 if i == 0 else 0)
        pf.space_after = Pt(6 if i == len(lines) - 1 else 0)
        pf.left_indent = Inches(0.25)
        pf.keep_together = True
        if i < len(lines) - 1:
            pf.keep_with_next = True
        shade(p._p.get_or_add_pPr(), CODE_BG)
        r = p.add_run(line if line.strip() else " ")
        r.font.name = MONO
        r.font.size = Pt(8.5)
        r.font.color.rgb = CODE_FG


def table(rows):
    """First row is the header; a |---| separator row is already stripped."""
    cols = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci in range(cols):
            cell = cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            text = row[ci] if ci < len(row) else ""
            if ri == 0:
                shade(cell._tc.get_or_add_tcPr(), "DEEAF6")
                run = p.add_run(re.sub(r"[*`]", "", text))
                run.bold = True
                run.font.size = Pt(9.5)
            else:
                inline(p, text)
                for r in p.runs:
                    r.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def split_row(line):
    """Split a pipe-table row into trimmed cells."""
    return [c.strip() for c in line.strip().strip("|").split("|")]


with open(SRC) as fh:
    lines = fh.read().split("\n")
i = 0
while i < len(lines):
    line = lines[i]

    # fenced code
    if line.startswith("```"):
        i += 1
        buf = []
        while i < len(lines) and not lines[i].startswith("```"):
            buf.append(lines[i])
            i += 1
        code_block(buf)
        i += 1
        continue

    # table: a header row followed by the |---| separator
    if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1]):
        rows = [split_row(line)]
        i += 2
        while i < len(lines) and lines[i].startswith("|"):
            rows.append(split_row(lines[i]))
            i += 1
        table(rows)
        continue

    stripped = line.strip()

    if stripped == "---":
        p = doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:color"), "BFBFBF")
        borders.append(bottom)
        pPr.append(borders)
        i += 1
        continue

    if stripped.startswith("# "):
        h = doc.add_heading(level=0)
        r = h.add_run(stripped[2:])
        r.font.color.rgb = ACCENT
        r.font.size = Pt(22)
    elif stripped.startswith("## "):
        h = doc.add_heading(level=1)
        h.paragraph_format.space_before = Pt(16)
        r = h.add_run(re.sub(r"[*`]", "", stripped[3:]))
        r.font.color.rgb = ACCENT
        r.font.size = Pt(15)
    elif stripped.startswith("### "):
        h = doc.add_heading(level=2)
        h.paragraph_format.space_before = Pt(10)
        r = h.add_run(re.sub(r"[*`]", "", stripped[4:]))
        r.font.color.rgb = ACCENT
        r.font.size = Pt(12)
    elif re.match(r"^[-*] ", stripped):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        inline(p, stripped[2:])
    elif re.match(r"^\d+\. ", stripped):
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        inline(p, re.sub(r"^\d+\.\s*", "", stripped))
    elif stripped == "":
        pass
    else:
        # Join wrapped prose: markdown hard-wraps, Word should reflow.
        buf = [stripped]
        while (
            i + 1 < len(lines)
            and lines[i + 1].strip()
            and not re.match(r"^([-*#>|]|\d+\.|```)", lines[i + 1].strip())
            and lines[i + 1].strip() != "---"
        ):
            i += 1
            buf.append(lines[i].strip())
        inline(doc.add_paragraph(), " ".join(buf))
    i += 1

doc.save(OUT)
print(f"wrote {OUT}")
